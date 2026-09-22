from datetime import datetime, time
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

from django.db import models
from .models import ToolTransaction, MaterialTransaction, Material

TO_EMAILS = [
    'vasudevb@kensova.ae',
    'deshanw@kensova.ae',
]
CC_EMAILS = [
    'prasadn@kensova.ae',
    'deshanir@kensova.ae',
    'dinushid@kensova.ae',
    'anushkaa@kensova.ae',
    'gayans@kensova.ae',
    'arunaa@kensova.ae',
]


def dubai_day_range(day=None):
    day = day or timezone.localdate()
    start = timezone.make_aware(datetime.combine(day, time.min))
    end = timezone.make_aware(datetime.combine(day, time.max))
    return start, end


def collect_daily_data(day=None):
    day = day or timezone.localdate()
    start, end = dubai_day_range(day)
    tools_issued = list(ToolTransaction.objects.filter(issue_date__range=(start, end)).select_related('tool', 'staff', 'project').order_by('issue_date'))
    tools_returned = list(ToolTransaction.objects.filter(return_date__range=(start, end)).select_related('tool', 'staff', 'project').order_by('return_date'))
    materials_issued = list(MaterialTransaction.objects.filter(transaction_type='Issue', date__range=(start, end)).select_related('material', 'staff', 'project').order_by('date'))
    materials_received = list(MaterialTransaction.objects.filter(transaction_type='Receive', date__range=(start, end)).select_related('material', 'supplier').order_by('date'))
    materials_returned = list(MaterialTransaction.objects.filter(transaction_type='Return', date__range=(start, end)).select_related('material', 'staff', 'project').order_by('date'))
    pending_tools = list(ToolTransaction.objects.filter(status='Issued').select_related('tool', 'staff', 'project').order_by('-issue_date'))
    low_stock = list(Material.objects.filter(is_deleted=False, current_stock__lte=0).order_by('material_code'))
    low_stock_alert = list(Material.objects.filter(is_deleted=False, current_stock__gt=0, current_stock__lte=models.F('min_stock_alert')).order_by('material_code'))
    return {
        'day': day,
        'tools_issued': tools_issued,
        'tools_returned': tools_returned,
        'materials_issued': materials_issued,
        'materials_received': materials_received,
        'materials_returned': materials_returned,
        'pending_tools': pending_tools,
        'out_of_stock': low_stock,
        'low_stock': low_stock_alert,
    }


def _fmt_dt(value):
    return timezone.localtime(value).strftime('%d-%b-%Y %H:%M') if value else '—'


def build_daily_pdf(data):
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=12*mm, leftMargin=12*mm, topMargin=12*mm, bottomMargin=12*mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='KTitle', parent=styles['Title'], fontSize=19, leading=22, alignment=TA_CENTER, spaceAfter=4))
    styles.add(ParagraphStyle(name='KSub', parent=styles['Normal'], fontSize=8.5, textColor=colors.HexColor('#475569'), alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle(name='KSection', parent=styles['Heading2'], fontSize=11, leading=14, textColor=colors.HexColor('#0f172a'), spaceBefore=10, spaceAfter=6))
    styles.add(ParagraphStyle(name='KSmall', parent=styles['Normal'], fontSize=7.2, leading=9))
    styles.add(ParagraphStyle(name='KSmallBold', parent=styles['Normal'], fontSize=7.2, leading=9, fontName='Helvetica-Bold'))

    story = [
        Paragraph('KENSOVA INTERIORS', styles['KTitle']),
        Paragraph(f'DAILY STORES MANAGEMENT REPORT — {data["day"].strftime("%d %B %Y")}', styles['KSub']),
        Paragraph('This is a system auto-generated report from Kensova ERP. The information reflects transactions and stock records captured in the system for the reporting date.', styles['KSmall']),
        Spacer(1, 5),
    ]

    cards = [
        ['TOOLS ISSUED', str(len(data['tools_issued'])), 'TOOLS RETURNED', str(len(data['tools_returned']))],
        ['PENDING TOOLS', str(len(data['pending_tools'])), 'MATERIAL ISSUES', str(len(data['materials_issued']))],
        ['MATERIAL RECEIPTS', str(len(data['materials_received'])), 'MATERIAL RETURNS', str(len(data['materials_returned']))],
    ]
    t = Table(cards, colWidths=[40*mm, 25*mm, 40*mm, 25*mm])
    t.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#f8fafc')),
        ('BOX',(0,0),(-1,-1),0.6,colors.HexColor('#cbd5e1')),
        ('INNERGRID',(0,0),(-1,-1),0.4,colors.HexColor('#e2e8f0')),
        ('FONTNAME',(0,0),(-1,-1),'Helvetica-Bold'),
        ('FONTSIZE',(0,0),(-1,-1),8),
        ('ALIGN',(1,0),(1,-1),'CENTER'),('ALIGN',(3,0),(3,-1),'CENTER'),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),
    ]))
    story += [t, Spacer(1, 5)]

    def section(title, headers, rows):
        story.append(Paragraph(title, styles['KSection']))
        if not rows:
            story.append(Paragraph('No records for this section.', styles['KSmall']))
            return
        data_rows = [[Paragraph(str(h), styles['KSmallBold']) for h in headers]]
        for row in rows:
            data_rows.append([Paragraph(str(c if c not in (None, '') else '—'), styles['KSmall']) for c in row])
        table = Table(data_rows, colWidths=[(186*mm)/len(headers)] * len(headers), repeatRows=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e0f2fe')),
            ('TEXTCOLOR',(0,0),(-1,0),colors.HexColor('#0c4a6e')),
            ('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#cbd5e1')),
            ('VALIGN',(0,0),(-1,-1),'TOP'),
            ('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),
            ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),
        ]))
        story.append(table)

    section('1. TOOLS — ISSUED TODAY', ['Time','Tool','Description','Employee','Project','Issued By'], [
        [_fmt_dt(x.issue_date), x.tool.tool_no, x.tool.tool_name, x.staff.formatted_name() if x.staff else x.received_by, x.project.project_name if x.project else 'Workshop F20', x.issued_by]
        for x in data['tools_issued']
    ])
    section('2. TOOLS — RETURNED TODAY', ['Time','Tool','Employee','Project','Condition','Received By Store'], [
        [_fmt_dt(x.return_date), x.tool.tool_no, x.staff.formatted_name() if x.staff else x.received_by, x.project.project_name if x.project else 'Workshop F20', x.return_condition or '—', x.received_by_store or '—']
        for x in data['tools_returned']
    ])
    section('3. MATERIALS — ISSUED TODAY', ['Time','Material','Description','Qty','Unit','Employee','Project','Issued By'], [
        [_fmt_dt(x.date), x.material.material_code, x.material.description, x.quantity, x.material.unit.upper(), x.staff.formatted_name() if x.staff else x.received_by, x.project.project_name if x.project else 'Workshop F20', x.issued_by]
        for x in data['materials_issued']
    ])
    section('4. MATERIALS — RECEIVED TODAY', ['Time','Material','Description','Qty','Unit','Supplier','GRN / Ref','Delivery Note'], [
        [_fmt_dt(x.date), x.material.material_code, x.material.description, x.quantity, x.material.unit.upper(), x.supplier.company_name if x.supplier else 'Direct Store Purchase', x.reference_no or '—', x.delivery_note_no or '—']
        for x in data['materials_received']
    ])
    section('5. MATERIALS — RETURNED TODAY', ['Time','Material','Description','Qty','Unit','Employee','Project'], [
        [_fmt_dt(x.date), x.material.material_code, x.material.description, x.quantity, x.material.unit.upper(), x.staff.formatted_name() if x.staff else x.received_by, x.project.project_name if x.project else 'Workshop F20']
        for x in data['materials_returned']
    ])
    section('6. MATERIAL STOCK ALERTS — END OF DAY', ['Code','Description','Current Stock','Unit','Minimum Alert','Status'], [
        [m.material_code, m.description, m.current_stock, m.unit.upper(), m.min_stock_alert, 'OUT OF STOCK']
        for m in data['out_of_stock']
    ] + [
        [m.material_code, m.description, m.current_stock, m.unit.upper(), m.min_stock_alert, 'LOW STOCK']
        for m in data['low_stock']
    ])
    section('7. PENDING TOOLS — END OF DAY', ['Tool','Description','Employee','Project','Issue Date'], [
        [x.tool.tool_no, x.tool.tool_name, x.staff.formatted_name() if x.staff else x.received_by, x.project.project_name if x.project else 'Workshop F20', _fmt_dt(x.issue_date)]
        for x in data['pending_tools']
    ])

    doc.build(story)
    buf.seek(0)
    return buf


def build_daily_excel(data):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Daily Summary'
    ws.append(['KENSOVA INTERIORS', 'Daily Stores Management Report', data['day'].strftime('%d-%b-%Y')])
    ws.append(['System Note', 'This is a system auto-generated report from Kensova ERP.'])
    ws.append([])
    ws.append(['Metric','Count'])
    ws.append(['Tools Issued', len(data['tools_issued'])])
    ws.append(['Tools Returned', len(data['tools_returned'])])
    ws.append(['Pending Tools', len(data['pending_tools'])])
    ws.append(['Material Issues', len(data['materials_issued'])])
    ws.append(['Material Receipts', len(data['materials_received'])])
    ws.append(['Material Returns', len(data['materials_returned'])])
    ws.append([])

    def add_sheet(title, headers, rows):
        sh = wb.create_sheet(title)
        sh.append(headers)
        for row in rows:
            sh.append(row)
        for cell in sh[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='DDEBF7')
        sh.freeze_panes = 'A2'
        for col in sh.columns:
            width = min(max(12, max(len(str(c.value or '')) for c in col) + 2), 35)
            sh.column_dimensions[get_column_letter(col[0].column)].width = width
        return sh

    add_sheet('Tools Issued', ['Time','Tool No','Description','Employee','Project','Issued By'], [
        [_fmt_dt(x.issue_date), x.tool.tool_no, x.tool.tool_name, x.staff.formatted_name() if x.staff else x.received_by or '', x.project.project_name if x.project else 'Workshop F20', x.issued_by]
        for x in data['tools_issued']
    ])
    add_sheet('Tools Returned', ['Time','Tool No','Description','Employee','Project','Condition','Received By Store'], [
        [_fmt_dt(x.return_date), x.tool.tool_no, x.tool.tool_name, x.staff.formatted_name() if x.staff else x.received_by or '', x.project.project_name if x.project else 'Workshop F20', x.return_condition or '', x.received_by_store or '']
        for x in data['tools_returned']
    ])
    add_sheet('Material Issues', ['Time','Code','Description','Qty','Unit','Employee','Project','Issued By'], [
        [_fmt_dt(x.date), x.material.material_code, x.material.description, x.quantity, x.material.unit.upper(), x.staff.formatted_name() if x.staff else x.received_by or '', x.project.project_name if x.project else 'Workshop F20', x.issued_by]
        for x in data['materials_issued']
    ])
    add_sheet('Material Receipts', ['Time','Code','Description','Qty','Unit','Supplier','GRN / Ref','Delivery Note'], [
        [_fmt_dt(x.date), x.material.material_code, x.material.description, x.quantity, x.material.unit.upper(), x.supplier.company_name if x.supplier else 'Direct Store Purchase', x.reference_no or '', x.delivery_note_no or '']
        for x in data['materials_received']
    ])
    add_sheet('Material Returns', ['Time','Code','Description','Qty','Unit','Employee','Project'], [
        [_fmt_dt(x.date), x.material.material_code, x.material.description, x.quantity, x.material.unit.upper(), x.staff.formatted_name() if x.staff else x.received_by or '', x.project.project_name if x.project else 'Workshop F20']
        for x in data['materials_returned']
    ])
    add_sheet('Stock Alerts', ['Code','Description','Current Stock','Unit','Minimum Alert','Status'], [
        [m.material_code, m.description, m.current_stock, m.unit.upper(), m.min_stock_alert, 'OUT OF STOCK'] for m in data['out_of_stock']
    ] + [
        [m.material_code, m.description, m.current_stock, m.unit.upper(), m.min_stock_alert, 'LOW STOCK'] for m in data['low_stock']
    ])
    add_sheet('Pending Tools', ['Tool No','Description','Employee','Project','Issue Date'], [
        [x.tool.tool_no, x.tool.tool_name, x.staff.formatted_name() if x.staff else x.received_by or '', x.project.project_name if x.project else 'Workshop F20', _fmt_dt(x.issue_date)]
        for x in data['pending_tools']
    ])
    for cell in ws[1]: cell.font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 28; ws.column_dimensions['B'].width = 32; ws.column_dimensions['C'].width = 18
    out = BytesIO(); wb.save(out); out.seek(0); return out


def send_via_outlook(pdf_bytes, xlsx_bytes, report_date):
    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError('Outlook Desktop integration requires pywin32. Install the project requirements first.') from exc

    outlook = win32com.client.Dispatch('Outlook.Application')
    mail = outlook.CreateItem(0)
    mail.To = '; '.join(TO_EMAILS)
    mail.CC = '; '.join(CC_EMAILS)
    mail.Subject = f'Kensova Stores Daily Management Report — {report_date.strftime("%d %B %Y")}'
    mail.HTMLBody = f'''<html><body style="font-family:Segoe UI,Arial,sans-serif;color:#172033">
    <h2 style="margin-bottom:4px">KENSOVA INTERIORS</h2>
    <p style="margin-top:0;color:#64748b">Stores & Materials — Daily Management Report</p>
    <p>Please find attached the professional daily stores report for <b>{report_date.strftime('%d %B %Y')}</b>, covering tool movements, material movements and pending tools.</p>
    <p style="font-size:12px;color:#64748b">This report was generated automatically by Kensova ERP.</p>
    </body></html>'''

    import tempfile
    pdf_path = Path(tempfile.gettempdir()) / f'Kensova_Daily_Report_{report_date:%Y%m%d}.pdf'
    xlsx_path = Path(tempfile.gettempdir()) / f'Kensova_Daily_Report_{report_date:%Y%m%d}.xlsx'
    pdf_path.write_bytes(pdf_bytes.getvalue())
    xlsx_path.write_bytes(xlsx_bytes.getvalue())
    mail.Attachments.Add(str(pdf_path))
    mail.Attachments.Add(str(xlsx_path))
    mail.Send()
    return True
