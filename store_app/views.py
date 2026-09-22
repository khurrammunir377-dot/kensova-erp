import json, random, io, os, urllib.request, urllib.error, sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import timedelta
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_POST
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.core.paginator import Paginator
from django.db.models import Sum, Count, Q, F, Case, When, IntegerField, OuterRef, Subquery
from django.db import transaction
from django.utils import timezone
from .models import Tool, ToolTransaction, Material, MaterialTransaction, Project, Supplier, Staff, PurchaseRequest, GatePass, GatePassItem
from .gate_pass_service import render_gate_pass_pdf
from .daily_report_service import collect_daily_data, build_daily_pdf, build_daily_excel, send_via_outlook, TO_EMAILS



def _pr_receipt_index():
    """Build receipt totals keyed by (PR number, material code) without changing DB schema.
    Receive transactions linked from the GRN form store a compact PR:<number> marker in remarks.
    """
    index = {}
    receipts = MaterialTransaction.objects.filter(transaction_type='Receive', remarks__startswith='PR:').select_related('material').only(
        'quantity', 'date', 'remarks', 'material__material_code'
    )
    for rc in receipts:
        marker, _, free_remarks = (rc.remarks or '').partition('|')
        pr_no = marker[3:].strip() if marker.startswith('PR:') else ''
        if not pr_no:
            continue
        key = (pr_no, rc.material.material_code)
        rec = index.setdefault(key, {'qty': 0, 'date': None, 'remarks': ''})
        rec['qty'] += int(rc.quantity or 0)
        if rec['date'] is None or (rc.date and rc.date > rec['date']):
            rec['date'] = rc.date
            rec['remarks'] = free_remarks.strip()
    return index

def _build_pr_item_rows(pr_queryset):
    receipt_index = _pr_receipt_index()
    rows = []
    for pr in pr_queryset:
        try:
            items = json.loads(pr.items_json or '[]')
        except (TypeError, ValueError, json.JSONDecodeError):
            items = []
        for item in items:
            code = str(item.get('code') or '').strip()
            requested = int(item.get('reorder_qty') or 0)
            linked = receipt_index.get((pr.pr_number, code), {'qty': 0, 'date': None, 'remarks': ''})
            received = int(linked.get('qty') or 0)
            pending = max(requested - received, 0)
            if requested > 0 and pending == 0:
                item_status = 'Received'
            elif received > 0:
                item_status = 'Partial'
            else:
                item_status = 'Pending'
            rows.append({
                'pr_number': pr.pr_number,
                'pr_date': pr.created_date,
                'item_code': code,
                'item_desc': item.get('desc') or '',
                'unit': str(item.get('unit') or 'PCS').upper(),
                'pr_qty': requested,
                'rec_qty': received,
                'pending_qty': pending,
                'item_status': item_status,
                'date_received': linked.get('date'),
                'remarks': linked.get('remarks') or pr.notes or '',
                'generated_by': pr.generated_by,
                'pr_status': pr.status,
                'pr_notes': pr.notes or '',
                'pr_items_json': pr.items_json or '[]',
            })
    return rows

def user_login(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    if request.method == 'POST':
        u = request.POST.get('username')
        p = request.POST.get('password')
        user = authenticate(request, username=u, password=p)
        if user:
            login(request, user)
            return redirect('dashboard')
        messages.error(request, 'Access Denied: Invalid Staff ID or Password.')
    return render(request, 'store_app/login.html')

def user_logout(request):
    logout(request)
    return redirect('login')

@login_required
def dashboard(request):
    total_tools = Tool.objects.filter(is_deleted=False).count()
    issued_tools = Tool.objects.filter(status='Issued', is_deleted=False).count()
    available_tools = Tool.objects.filter(status='Available', is_deleted=False).count()
    maintenance_tools = Tool.objects.filter(status='Maintenance', is_deleted=False).count()
    
    materials = Material.objects.filter(is_deleted=False)
    total_materials = materials.count()
    low_stock_materials = materials.filter(current_stock__lte=F('min_stock_alert')).count()
    active_projects = Project.objects.filter(status='Active', is_deleted=False).count()
    total_staff = Staff.objects.filter(is_deleted=False).count()

    today = timezone.now().date()
    overdue_tools_count = ToolTransaction.objects.filter(status='Issued', expected_return_date__lt=today).count()

    tool_chart_labels = ['Available in F20', 'Issued to Sites', 'Under Maintenance']
    tool_chart_data = [available_tools, issued_tools, maintenance_tools]
    
    top_mats = materials.order_by('-current_stock')[:8]
    mat_labels = [m.material_code for m in top_mats]
    mat_data = [int(m.current_stock) for m in top_mats]

    site_alloc = ToolTransaction.objects.filter(status='Issued').values('project__project_name').annotate(total=Count('id')).order_by('-total')[:5]
    site_labels = [s['project__project_name'] or 'Workshop F20' for s in site_alloc] or ['No Active Sites']
    site_data = [s['total'] for s in site_alloc] or [0]

    days_7_labels = []
    inward_7_trend = []
    outward_7_trend = []
    for i in range(6, -1, -1):
        target_day = today - timedelta(days=i)
        days_7_labels.append(target_day.strftime('%d-%b'))
        in_qty = MaterialTransaction.objects.filter(transaction_type='Receive', date__date=target_day).aggregate(s=Sum('quantity'))['s'] or 0
        out_qty = MaterialTransaction.objects.filter(transaction_type='Issue', date__date=target_day).aggregate(s=Sum('quantity'))['s'] or 0
        inward_7_trend.append(int(in_qty))
        outward_7_trend.append(int(out_qty))

    past_30_days_labels = []
    daily_tool_issues = []
    daily_tool_returns = []
    daily_mat_issues = []
    daily_mat_receipts = []

    for i in range(29, -1, -1):
        target_day = today - timedelta(days=i)
        past_30_days_labels.append(target_day.strftime('%d-%b'))
        
        t_iss = ToolTransaction.objects.filter(issue_date__date=target_day).count()
        t_ret = ToolTransaction.objects.filter(return_date__date=target_day, status='Returned').count()
        daily_tool_issues.append(t_iss)
        daily_tool_returns.append(t_ret)

        m_iss = MaterialTransaction.objects.filter(transaction_type='Issue', date__date=target_day).aggregate(s=Sum('quantity'))['s'] or 0
        m_rec = MaterialTransaction.objects.filter(transaction_type='Receive', date__date=target_day).aggregate(s=Sum('quantity'))['s'] or 0
        daily_mat_issues.append(int(m_iss))
        daily_mat_receipts.append(int(m_rec))

    context = {
        'page_title': 'System Dashboard',
        'page_subtitle': 'Real-time telemetry of tools, materials, sites, and store movements',
        'page_icon': 'fa-solid fa-chart-pie',
        'page_icon_color': 'from-emerald-500 to-teal-600',
        'total_tools': total_tools,
        'issued_tools': issued_tools,
        'available_tools': available_tools,
        'maintenance_tools': maintenance_tools,
        'total_materials': total_materials,
        'low_stock_materials': low_stock_materials,
        'active_projects': active_projects,
        'total_staff': total_staff,
        'overdue_tools_count': overdue_tools_count,
        'tool_chart_labels_json': json.dumps(tool_chart_labels),
        'tool_chart_data_json': json.dumps(tool_chart_data),
        'mat_labels_json': json.dumps(mat_labels),
        'mat_data_json': json.dumps(mat_data),
        'site_labels_json': json.dumps(site_labels),
        'site_data_json': json.dumps(site_data),
        'days_7_labels_json': json.dumps(days_7_labels),
        'inward_7_trend_json': json.dumps(inward_7_trend),
        'outward_7_trend_json': json.dumps(outward_7_trend),
        'past_30_days_labels_json': json.dumps(past_30_days_labels),
        'daily_tool_issues_json': json.dumps(daily_tool_issues),
        'daily_tool_returns_json': json.dumps(daily_tool_returns),
        'daily_mat_issues_json': json.dumps(daily_mat_issues),
        'daily_mat_receipts_json': json.dumps(daily_mat_receipts),
    }
    return render(request, 'store_app/dashboard.html', context)

@login_required
def tool_issue(request):
    custodian_name = f"{request.user.username}, {request.user.get_full_name() or 'Manuja Shehan'}"
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'quick_add_tool':
            t_no = request.POST.get('tool_no').strip()
            t_name = request.POST.get('tool_name').strip()
            Tool.objects.create(
                tool_no=t_no,
                tool_name=t_name,
                store_location=request.POST.get('store_location', 'F20'),
                status=request.POST.get('status', 'Available'),
                calibration_due_date=request.POST.get('calibration_due_date') or None
            )
            messages.success(request, f"✓ Added Tool [{t_no}] to Store F20 catalog.")
            return redirect('tool_issue')

        tool = get_object_or_404(Tool, tool_no=request.POST.get('tool_no'))
        staff = get_object_or_404(Staff, staff_no=request.POST.get('staff_no'))
        project_code = request.POST.get('project_code')
        project = Project.objects.filter(project_code=project_code).first()
        issued_by_val = request.POST.get('issued_by', custodian_name)
        gate_pass = f"GP-DXB-{timezone.now().strftime('%y%m')}-{random.randint(1000, 9999)}"

        ToolTransaction.objects.create(
            tool=tool,
            staff=staff,
            project=project,
            issued_by=issued_by_val,
            received_by=staff.formatted_name(),
            gate_pass_no=gate_pass,
            remarks=request.POST.get('remarks')
        )
        tool.status = 'Issued'
        tool.status_changed_at = timezone.now()
        tool.save(update_fields=['status', 'status_changed_at'])
        messages.success(request, f"✓ Issued [{tool.tool_no}] ({tool.tool_name}) to {staff.formatted_name()}.")
        return redirect(f"/tools/issue/?selected_staff={staff.staff_no}&selected_project={project_code or ''}")

    selected_staff = request.GET.get('selected_staff', '')
    selected_project = request.GET.get('selected_project', '')

    available_tools = Tool.objects.filter(status='Available', is_deleted=False).order_by('tool_no')
    staff_list = Staff.objects.filter(is_deleted=False).order_by('staff_no')
    recent_issues = ToolTransaction.objects.filter(status='Issued').select_related('tool', 'staff', 'project')[:9]
    pending_count = ToolTransaction.objects.filter(status='Issued').count()

    return render(request, 'store_app/tool_issue.html', {
        'page_title': 'Issue Tool to Technician',
        'page_subtitle': 'Store Handover Form • Location F20 Checkout',
        'page_icon': 'fa-solid fa-arrow-up-right-from-square',
        'page_icon_color': 'from-amber-500 to-orange-600',
        'current_custodian': custodian_name,
        'available_tools': available_tools,
        'staff_list': staff_list,
        'projects': Project.objects.filter(status='Active', is_deleted=False),
        'recent_issues': recent_issues,
        'pending_count': pending_count,
        'selected_staff': selected_staff,
        'selected_project': selected_project,
    })

@login_required
def tool_return(request):
    custodian_name = f"{request.user.username}, {request.user.get_full_name() or 'Manuja Shehan'}"
    if request.method == 'POST':
        tx = get_object_or_404(ToolTransaction, id=request.POST.get('tx_id'), status='Issued')
        cond = request.POST.get('return_condition', 'Good')
        returned_by_val = request.POST.get('returned_by', tx.staff.formatted_name())
        received_by_store_val = request.POST.get('received_by_store', custodian_name)
        remarks = request.POST.get('remarks', '')

        tx.status = 'Returned'
        tx.return_date = timezone.now()
        tx.return_condition = cond
        tx.received_by_store = received_by_store_val
        if remarks:
            tx.remarks = f"{tx.remarks or ''} | Inspection: {remarks} (Handed by: {returned_by_val})"
        tx.save()

        tool = tx.tool
        tool.status = 'Maintenance' if cond in ['Damaged', 'Needs Repair'] else 'Available'
        tool.status_changed_at = timezone.now()
        tool.save(update_fields=['status', 'status_changed_at'])

        messages.success(request, f"✓ Returned [{tool.tool_no}] ({tool.tool_name}) into Store F20 ({cond}).")
        return redirect('tool_return')

    open_txs = ToolTransaction.objects.filter(status='Issued').select_related('tool', 'staff', 'project').order_by('tool__tool_no', '-issue_date')
    pending_txs = []
    seen_tools = set()
    for open_tx in open_txs:
        if open_tx.tool_id not in seen_tools:
            pending_txs.append(open_tx)
            seen_tools.add(open_tx.tool_id)
    recent_returns = ToolTransaction.objects.filter(status='Returned').select_related('tool', 'staff')[:9]
    pending_count = ToolTransaction.objects.filter(status='Issued').count()
    selected_tx = request.GET.get('tx_id', '')
    selected_returned_by = next(
        (tx.staff.formatted_name() for tx in pending_txs if str(tx.id) == str(selected_tx)),
        ''
    )

    return render(request, 'store_app/tool_return.html', {
        'page_title': 'Process Tool Return',
        'page_subtitle': 'Store Intake Form • Location F20 Restock',
        'page_icon': 'fa-solid fa-arrow-rotate-left',
        'page_icon_color': 'from-emerald-500 to-teal-600',
        'current_custodian': custodian_name,
        'pending_txs': pending_txs,
        'recent_returns': recent_returns,
        'pending_count': pending_count,
        'selected_tx': selected_tx,
        'selected_returned_by': selected_returned_by,
    })

@login_required
def tool_pending(request):
    if request.method == 'POST' and request.POST.get('action') == 'generate_gate_pass':
        selected_ids = [v for v in request.POST.getlist('selected_tx') if str(v).isdigit()]
        return _create_gate_pass_from_tool_ids(request, selected_ids)

    if request.method == 'POST' and request.POST.get('action') == 'bulk_return':
        selected_ids = [v for v in request.POST.getlist('selected_tx') if str(v).isdigit()]
        if not selected_ids:
            messages.error(request, 'Select at least one pending tool to return.')
            return redirect('tool_pending')
        custodian_name = f"{request.user.username}, {request.user.get_full_name() or 'Manuja Shehan'}"
        txs = ToolTransaction.objects.filter(id__in=selected_ids, status='Issued').select_related('tool')
        returned = 0
        now = timezone.now()
        for tx in txs:
            tx.status = 'Returned'
            tx.return_date = now
            tx.return_condition = 'Good'
            tx.received_by_store = custodian_name
            tx.remarks = f"{tx.remarks or ''} | Bulk return confirmed from Pending Tools".strip()
            tx.save()
            tool = tx.tool
            tool.status = 'Available'
            tool.status_changed_at = now
            tool.save(update_fields=['status', 'status_changed_at'])
            returned += 1
        messages.success(request, f'✓ {returned} selected tool(s) returned to Store F20.')
        return redirect('tool_pending')

    q = request.GET.get('q', '')
    filter_site = request.GET.get('project', '')
    pending_list = ToolTransaction.objects.filter(status='Issued').select_related('tool', 'staff', 'project')
    
    if q:
        pending_list = pending_list.filter(
            Q(tool__tool_no__icontains=q) |
            Q(tool__tool_name__icontains=q) |
            Q(staff__name__icontains=q) |
            Q(staff__staff_no__icontains=q) |
            Q(issued_by__icontains=q) |
            Q(remarks__icontains=q)
        )
    if filter_site:
        pending_list = pending_list.filter(project__project_code=filter_site)

    filtered_count = pending_list.count()

    return render(request, 'store_app/tool_pending.html', {
        'page_title': 'Pending Tools on Active Sites',
        'page_subtitle': 'Real-time monitoring of checked-out equipment across projects',
        'page_icon': 'fa-solid fa-clock-rotate-left',
        'page_icon_color': 'from-rose-500 to-red-600',
        'pending_list': pending_list,
        'filtered_count': filtered_count,
        'q': q,
        'projects': Project.objects.filter(status='Active', is_deleted=False),
        'selected_project': filter_site,
        'today': timezone.now().date(),
    })

@login_required
def material_inventory(request):
    # Legacy catalogue route retained for compatibility; the active workspace is Records & Reports.
    return redirect('/reports/?tab=materialstock')

@login_required
def material_issue(request):
    custodian_name = f"{request.user.username}, {request.user.get_full_name() or 'Manuja Shehan'}"
    if request.method == 'POST':
        mat = get_object_or_404(Material, material_code=request.POST.get('material_code'))
        qty = int(request.POST.get('quantity', 1))
        project_code = request.POST.get('project_code')
        staff_no = request.POST.get('staff_no')
        staff_obj = Staff.objects.filter(staff_no=staff_no).first()

        if int(mat.current_stock) < qty:
            messages.error(request, f"Cannot issue! Only {mat.current_stock} {mat.unit.upper()} remaining in Store F20.")
        else:
            mat.current_stock = int(mat.current_stock) - qty
            mat.save()
            MaterialTransaction.objects.create(
                material=mat,
                transaction_type='Issue',
                quantity=qty,
                project=Project.objects.filter(project_code=project_code).first(),
                staff=staff_obj,
                issued_by=request.POST.get('issued_by', custodian_name),
                received_by=staff_obj.formatted_name() if staff_obj else '',
                reference_no=request.POST.get('reference_no'),
                remarks=request.POST.get('remarks')
            )
            messages.success(request, f"✓ Issued {qty} {mat.unit.upper()} of {mat.description}. (Remaining: {mat.current_stock} {mat.unit.upper()}).")
        return redirect(f"/materials/issue/?selected_staff={staff_no or ''}&selected_project={project_code or ''}")

    selected_staff = request.GET.get('selected_staff', '')
    selected_project = request.GET.get('selected_project', '')

    materials = Material.objects.filter(current_stock__gt=0, is_deleted=False).order_by('material_code')
    staff_list = Staff.objects.filter(is_deleted=False).order_by('staff_no')
    recent_issues = MaterialTransaction.objects.filter(transaction_type='Issue').select_related('material', 'project')[:9]

    return render(request, 'store_app/material_issue.html', {
        'page_title': 'Issue Material to Project Site',
        'page_subtitle': 'Direct warehouse stock deduction • Location F20 Outward',
        'page_icon': 'fa-solid fa-dolly',
        'page_icon_color': 'from-purple-600 to-pink-600',
        'current_custodian': custodian_name,
        'materials': materials,
        'projects': Project.objects.filter(status='Active', is_deleted=False),
        'staff_list': staff_list,
        'recent_issues': recent_issues,
        'selected_staff': selected_staff,
        'selected_project': selected_project,
    })

@login_required
def material_receive(request):
    custodian_name = f"{request.user.username}, {request.user.get_full_name() or 'Manuja Shehan'}"
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'quick_add_material':
            m_code = request.POST.get('material_code').strip()
            m_desc = request.POST.get('description').strip()
            Material.objects.create(
                material_code=m_code,
                description=m_desc,
                specification=request.POST.get('specification', ''),
                current_stock=int(request.POST.get('current_stock', 0)),
                unit=request.POST.get('unit', 'pcs'),
                min_stock_alert=20,
                store_location='F20'
            )
            messages.success(request, f"✓ Material [{m_code}] added to catalog.")
            return redirect('material_receive')

        mat = get_object_or_404(Material, material_code=request.POST.get('material_code'))
        pr_number = (request.POST.get('pr_number') or '').strip()
        if pr_number:
            linked_pr = PurchaseRequest.objects.filter(pr_number=pr_number).first()
            if not linked_pr:
                messages.error(request, f"PR [{pr_number}] was not found. Select a valid PR number or leave the field blank.")
                return redirect('material_receive')
            try:
                linked_items = json.loads(linked_pr.items_json or '[]')
            except (TypeError, ValueError, json.JSONDecodeError):
                linked_items = []
            if mat.material_code not in {str(i.get('code') or '') for i in linked_items}:
                messages.error(request, f"Material [{mat.material_code}] is not listed on PR [{pr_number}].")
                return redirect('material_receive')

        qty = int(request.POST.get('quantity', 1))
        mat.current_stock = int(mat.current_stock) + qty
        mat.save()

        supplier_code = request.POST.get('supplier_code')
        grn_ref = request.POST.get('grn_number')
        delivery_note = request.POST.get('delivery_note_no') or request.POST.get('reference_no')
        pr_marker = f"PR:{pr_number}" if pr_number else ''
        free_remarks = (request.POST.get('remarks') or '').strip()
        linked_remarks = f"{pr_marker}|{free_remarks}" if pr_marker else free_remarks
        MaterialTransaction.objects.create(
            material=mat,
            transaction_type='Receive',
            quantity=qty,
            supplier=Supplier.objects.filter(supplier_code=supplier_code).first(),
            issued_by=request.POST.get('received_by_store', custodian_name),
            received_by=request.POST.get('delivered_by', ''),
            reference_no=grn_ref,
            delivery_note_no=delivery_note,
            remarks=linked_remarks
        )
        messages.success(request, f"✓ GRN [{grn_ref}]: +{qty} {mat.unit.upper()} received for {mat.description} (Stock: {mat.current_stock} {mat.unit.upper()}).")
        return redirect('material_receive')

    recent_grns = MaterialTransaction.objects.filter(transaction_type='Receive').select_related('material', 'supplier')[:9]

    return render(request, 'store_app/material_receive.html', {
        'page_title': 'Receive Goods (GRN Inward)',
        'page_subtitle': 'Goods Received Note • Location F20 Inward Form',
        'page_icon': 'fa-solid fa-truck-ramp-box',
        'page_icon_color': 'from-teal-500 to-cyan-600',
        'current_custodian': custodian_name,
        'materials': Material.objects.filter(is_deleted=False).order_by('material_code'),
        'suppliers': Supplier.objects.filter(is_deleted=False).order_by('company_name'),
        'recent_grns': recent_grns,
        'open_prs': PurchaseRequest.objects.all().order_by('-created_date')[:100],
    })

@login_required
def material_return(request):
    custodian_name = f"{request.user.username}, {request.user.get_full_name() or 'Manuja Shehan'}"
    if request.method == 'POST':
        mat = get_object_or_404(Material, material_code=request.POST.get('material_code'))
        qty = int(request.POST.get('quantity', 1))
        mat.current_stock = int(mat.current_stock) + qty
        mat.save()

        project_code = request.POST.get('project_code')
        staff_no = request.POST.get('staff_no')
        staff_obj = Staff.objects.filter(staff_no=staff_no).first()

        MaterialTransaction.objects.create(
            material=mat,
            transaction_type='Return',
            quantity=qty,
            project=Project.objects.filter(project_code=project_code).first(),
            staff=staff_obj,
            issued_by=request.POST.get('received_by_store', custodian_name),
            received_by=request.POST.get('returned_by', (staff_obj.formatted_name() if staff_obj else '')),
            remarks=request.POST.get('remarks')
        )
        messages.success(request, f"✓ Returned {qty} {mat.unit.upper()} of {mat.description} into Store F20.")
        return redirect(f"/materials/return/?selected_staff={staff_no or ''}&selected_project={project_code or ''}")

    selected_staff = request.GET.get('selected_staff', '')
    selected_project = request.GET.get('selected_project', '')

    recent_returns = MaterialTransaction.objects.filter(transaction_type='Return').select_related('material', 'project')[:9]
    staff_list = Staff.objects.filter(is_deleted=False).order_by('staff_no')

    return render(request, 'store_app/material_return.html', {
        'page_title': 'Return Surplus Material from Site',
        'page_subtitle': 'Restock leftover site consumables into Store F20 Form',
        'page_icon': 'fa-solid fa-rotate-left',
        'page_icon_color': 'from-indigo-500 to-blue-600',
        'current_custodian': custodian_name,
        'materials': Material.objects.filter(is_deleted=False).order_by('material_code'),
        'projects': Project.objects.filter(is_deleted=False),
        'staff_list': staff_list,
        'recent_returns': recent_returns,
        'selected_staff': selected_staff,
        'selected_project': selected_project,
    })

@login_required
def projects_list(request):
    # Legacy sites route retained for compatibility; the active workspace is Records & Reports.
    return redirect('/reports/?tab=projects')

@login_required
def suppliers_list(request):
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'create':
            Supplier.objects.create(
                supplier_code=request.POST.get('supplier_code').strip(),
                company_name=request.POST.get('company_name').strip(),
                contact_person=request.POST.get('contact_person'),
                phone=request.POST.get('phone'),
                email=request.POST.get('email'),
                address=request.POST.get('address')
            )
            messages.success(request, "✓ Supplier registered in vendor directory.")
        elif action == 'edit':
            s = get_object_or_404(Supplier, supplier_code=request.POST.get('supplier_code'))
            s.company_name = request.POST.get('company_name')
            s.contact_person = request.POST.get('contact_person')
            s.phone = request.POST.get('phone')
            s.email = request.POST.get('email')
            s.address = request.POST.get('address')
            s.save()
            messages.success(request, f"✓ Supplier [{s.supplier_code}] updated.")
        elif action == 'delete':
            s = get_object_or_404(Supplier, supplier_code=request.POST.get('supplier_code'))
            s.is_deleted = True
            s.save()
            messages.success(request, f"✓ Supplier [{s.supplier_code}] removed.")
        return redirect('suppliers_list')

    suppliers = Supplier.objects.filter(is_deleted=False).order_by('company_name')
    return render(request, 'store_app/suppliers_list.html', {
        'page_title': 'Suppliers & Vendor Directory',
        'page_subtitle': 'Procurement & Store Materials Supply Partners',
        'page_icon': 'fa-solid fa-building-user',
        'page_icon_color': 'from-blue-500 to-indigo-600',
        'suppliers': suppliers,
    })

@login_required
def export_reports_excel(request):
    tab = request.GET.get('tab', 'toolhistory')
    q = request.GET.get('q', '')
    state = request.GET.get('state', 'all')
    status = request.GET.get('status', '')
    stock = request.GET.get('stock', '')
    project_filter = request.GET.get('project', '')
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tab[:31]

    header_font = Font(name='Arial', size=11, bold=True, color='0F172A')
    header_fill = PatternFill(start_color='E0F2FE', end_color='E0F2FE', fill_type='solid')
    border_thin = Border(left=Side(style='thin', color='CBD5E1'),
                         right=Side(style='thin', color='CBD5E1'),
                         top=Side(style='thin', color='CBD5E1'),
                         bottom=Side(style='thin', color='CBD5E1'))

    if tab in ('gatepassall', 'materialgp', 'toolsgp'):
        headers = ['Pass No', 'Type', 'Created At', 'Employee', 'Employee ID / Designation', 'Destination', 'Item Count', 'Created By']
        ws.append(headers)
        qs = GatePass.objects.all().prefetch_related('items').order_by('-created_at')
        if tab == 'materialgp':
            qs = qs.filter(pass_type='Material')
        elif tab == 'toolsgp':
            qs = qs.filter(pass_type='Tool')
        if q:
            qs = qs.filter(Q(movement_pass_no__icontains=q) | Q(employee_name__icontains=q) | Q(destination__icontains=q))
        if date_from:
            qs = qs.filter(created_at__date__gte=date_from)
        if date_to:
            qs = qs.filter(created_at__date__lte=date_to)
        for gp in qs:
            ws.append([gp.movement_pass_no, gp.get_pass_type_display(), gp.created_at.strftime('%d-%b-%Y %H:%M'), gp.employee_name or '—', gp.employee_id_designation or '—', gp.destination or '—', gp.items.count(), gp.created_by or '—'])

    elif tab == 'toolhistory':
        headers = ['Tool No', 'Tool Name', 'Issued At', 'Issued By', 'Holder / Technician', 'Project', 'Returned At', 'Received By Store', 'Condition']
        ws.append(headers)
        for r in ToolTransaction.objects.select_related('tool', 'staff', 'project').order_by('-issue_date'):
            ws.append([
                r.tool.tool_no, r.tool.tool_name,
                r.issue_date.strftime('%d-%b-%Y %H:%M') if r.issue_date else '',
                r.issued_by,
                r.staff.formatted_name() if r.staff else r.received_by,
                r.project.project_name if r.project else 'Workshop F20',
                r.return_date.strftime('%d-%b-%Y %H:%M') if r.return_date else 'PENDING',
                r.received_by_store or '—',
                r.return_condition or '—'
            ])
    elif tab == 'toolregister':
        headers = ['Tool No', 'Tool Name', 'Store Location', 'Status', 'Status Date', 'Calibration Due Date']
        ws.append(headers)
        for t in Tool.objects.filter(is_deleted=False).order_by('tool_no'):
            ws.append([t.tool_no, t.tool_name, t.store_location, t.status,
                       t.status_changed_at.strftime('%d-%b-%Y %H:%M') if t.status_changed_at else '',
                       str(t.calibration_due_date or 'Not Set')])
    elif tab == 'materialstock':
        headers = ['Part No', 'Description / Specification', 'Qty', 'Unit', 'Minimum', 'Location', 'GRN Details']
        ws.append(headers)
        for m in Material.objects.filter(is_deleted=False).order_by('material_code'):
            grn = MaterialTransaction.objects.filter(transaction_type='Receive', material=m).order_by('-date').first()
            grn_details = f"{grn.reference_no or '—'} / {grn.delivery_note_no or '—'}" if grn else '—'
            desc = m.description + (f"\n{m.specification}" if m.specification else '')
            ws.append([m.material_code, desc, int(m.current_stock), m.unit.upper(), int(m.min_stock_alert), m.store_location, grn_details])
    elif tab == 'issues':
        headers = ['Date', 'Material Code', 'Description', 'Qty', 'Unit', 'Issued By', 'Received By', 'Project', 'Reference']
        ws.append(headers)
        for i in MaterialTransaction.objects.filter(transaction_type='Issue').select_related('material', 'staff', 'project').order_by('-date'):
            ws.append([
                i.date.strftime('%d-%b-%Y %H:%M') if i.date else '',
                i.material.material_code, i.material.description, int(i.quantity), i.material.unit.upper(),
                i.issued_by, (i.received_by or (i.staff.formatted_name() if i.staff else '—')),
                i.project.project_name if i.project else 'Workshop', i.reference_no or '—'
            ])
    elif tab == 'receipts':
        headers = ['Date', 'Material Code', 'Description', 'Qty', 'Unit', 'Supplier', 'GRN No', 'Delivery Note / PO', 'Delivered By', 'Received By Store']
        ws.append(headers)
        for rc in MaterialTransaction.objects.filter(transaction_type='Receive').select_related('material', 'supplier').order_by('-date'):
            ws.append([
                rc.date.strftime('%d-%b-%Y %H:%M') if rc.date else '',
                rc.material.material_code, rc.material.description, int(rc.quantity), rc.material.unit.upper(),
                rc.supplier.company_name if rc.supplier else 'Direct Purchase', rc.reference_no or '—',
                rc.delivery_note_no or '—', rc.received_by or '—', rc.issued_by
            ])
    elif tab == 'materialreturns':
        headers = ['Date', 'Material Code', 'Description', 'Qty', 'Unit', 'From Project', 'Returned By', 'Received By Store', 'Reference']
        ws.append(headers)
        for rt in MaterialTransaction.objects.filter(transaction_type='Return').select_related('material', 'staff', 'project').order_by('-date'):
            ws.append([
                rt.date.strftime('%d-%b-%Y %H:%M') if rt.date else '',
                rt.material.material_code, rt.material.description, int(rt.quantity), rt.material.unit.upper(),
                rt.project.project_name if rt.project else 'Site', (rt.received_by or (rt.staff.formatted_name() if rt.staff else '—')),
                rt.issued_by, rt.reference_no or '—'
            ])
    elif tab == 'purchase':
        headers = ['PR No', 'PR Date', 'Item Detail', 'Unit', 'PR Qty', 'Rec Qty', 'Pending Qty', 'Status', 'Date Rec', 'Remarks']
        ws.append(headers)
        for item in _build_pr_item_rows(PurchaseRequest.objects.all().order_by('-created_date')):
            ws.append([
                item['pr_number'],
                item['pr_date'].strftime('%d-%b-%Y') if item['pr_date'] else '',
                f"[{item['item_code']}] {item['item_desc']}", item['unit'], item['pr_qty'], item['rec_qty'], item['pending_qty'],
                item['item_status'], item['date_received'].strftime('%d-%b-%Y') if item['date_received'] else '', item['remarks']
            ])
    elif tab == 'staff':
        headers = ['Staff ID', 'Full Name', 'Designation', 'Store Custodian']
        ws.append(headers)
        for s in Staff.objects.filter(is_deleted=False).order_by('staff_no'):
            ws.append([s.staff_no, s.name, s.designation, 'Yes' if s.is_store_custodian else 'No'])
    elif tab == 'suppliers':
        headers = ['Code', 'Company Name', 'Contact Person', 'Phone', 'Email', 'Address']
        ws.append(headers)
        for sup in Supplier.objects.filter(is_deleted=False).order_by('company_name'):
            ws.append([sup.supplier_code, sup.company_name, sup.contact_person or '', sup.phone or '', sup.email or '', sup.address or ''])
    elif tab == 'projects':
        headers = ['Code', 'Project Name', 'Client Name', 'Location', 'Budget AED', 'Start Date', 'Completion Date', 'Status']
        ws.append(headers)
        for p in Project.objects.filter(is_deleted=False).order_by('-start_date'):
            ws.append([p.project_code, p.project_name, p.client_name, p.location, float(p.budget_aed), str(p.start_date), str(p.completion_date or 'TBD'), p.status])
    else:
        headers = ['Deleted Time', 'Record Type', 'Identifier', 'Label']
        ws.append(headers)
        for t in Tool.objects.filter(is_deleted=True): ws.append([str(timezone.now().date()), 'Tool', t.tool_no, t.tool_name])
        for m in Material.objects.filter(is_deleted=True): ws.append([str(timezone.now().date()), 'Material', m.material_code, m.description])
        for p in Project.objects.filter(is_deleted=True): ws.append([str(timezone.now().date()), 'Project', p.project_code, p.project_name])
        for s in Supplier.objects.filter(is_deleted=True): ws.append([str(timezone.now().date()), 'Supplier', s.supplier_code, s.company_name])
        for st in Staff.objects.filter(is_deleted=True): ws.append([str(timezone.now().date()), 'Staff', st.staff_no, st.name])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border_thin

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Kensova_Export_{tab}_{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response = HttpResponse(output.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response

@login_required
def reports_hub(request):
    tab = request.GET.get('tab', 'toolhistory')
    q = request.GET.get('q', '')
    state = request.GET.get('state', 'all')
    status = request.GET.get('status', '')
    stock = request.GET.get('stock', '')
    project_filter = request.GET.get('project', '')
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'generate_material_gate_pass':
            selected_ids = [v for v in request.POST.getlist('selected_material_tx') if str(v).isdigit()]
            return _create_gate_pass_from_material_ids(request, selected_ids)
        if action == 'add_tool':
            Tool.objects.create(
                tool_no=request.POST.get('tool_no').strip(),
                tool_name=request.POST.get('tool_name').strip(),
                store_location=request.POST.get('store_location', 'F20').strip(),
                status=request.POST.get('status', 'Available'),
                calibration_due_date=request.POST.get('calibration_due_date') or None
            )
            messages.success(request, f"✓ Tool [{request.POST.get('tool_no')}] registered successfully.")
            return redirect('/reports/?tab=toolregister')
        elif action == 'edit_tool':
            tool = get_object_or_404(Tool, tool_no=request.POST.get('tool_no'), is_deleted=False)
            new_status = request.POST.get('status', tool.status)
            tool.tool_name = request.POST.get('tool_name', tool.tool_name).strip()
            tool.store_location = request.POST.get('store_location', tool.store_location).strip()
            tool.serial_no = request.POST.get('serial_no') or None
            tool.calibration_due_date = request.POST.get('calibration_due_date') or None
            if new_status != tool.status:
                tool.status = new_status
                tool.status_changed_at = timezone.now()
            tool.save()
            messages.success(request, f"✓ Tool [{tool.tool_no}] updated.")
            return redirect('/reports/?tab=toolregister')
        elif action == 'delete_tool':
            tool = get_object_or_404(Tool, tool_no=request.POST.get('tool_no'), is_deleted=False)
            if tool.status == 'Issued':
                messages.error(request, 'Issued tools cannot be deleted. Return the tool first.')
            else:
                tool.is_deleted = True
                tool.save(update_fields=['is_deleted'])
                messages.success(request, f"✓ Tool [{tool.tool_no}] removed.")
            return redirect('/reports/?tab=toolregister')
        elif action in ('edit_material', 'delete_material'):
            material = get_object_or_404(Material, material_code=request.POST.get('material_code'), is_deleted=False)
            if action == 'delete_material':
                material.is_deleted = True
                material.save(update_fields=['is_deleted'])
                messages.success(request, f"✓ Material [{material.material_code}] removed.")
            else:
                material.description = request.POST.get('description', material.description).strip()
                material.specification = request.POST.get('specification', '')
                material.current_stock = int(request.POST.get('current_stock', material.current_stock))
                material.unit = request.POST.get('unit', material.unit).strip()
                material.min_stock_alert = int(request.POST.get('min_stock_alert', material.min_stock_alert))
                material.store_location = request.POST.get('store_location', material.store_location).strip()
                material.save()
                messages.success(request, f"✓ Material [{material.material_code}] updated.")
            return redirect('/reports/?tab=materialstock')
        elif action in ('edit_supplier', 'delete_supplier'):
            supplier = get_object_or_404(Supplier, supplier_code=request.POST.get('supplier_code'), is_deleted=False)
            if action == 'delete_supplier':
                supplier.is_deleted = True
                supplier.save(update_fields=['is_deleted'])
                messages.success(request, f"✓ Supplier [{supplier.supplier_code}] removed.")
            else:
                supplier.company_name = request.POST.get('company_name', supplier.company_name).strip()
                supplier.contact_person = request.POST.get('contact_person', '')
                supplier.phone = request.POST.get('phone', '')
                supplier.email = request.POST.get('email') or None
                supplier.address = request.POST.get('address', '')
                supplier.save()
                messages.success(request, f"✓ Supplier [{supplier.supplier_code}] updated.")
            return redirect('/reports/?tab=suppliers')
        elif action == 'create_supplier':
            Supplier.objects.create(
                supplier_code=request.POST.get('supplier_code', '').strip(),
                company_name=request.POST.get('company_name', '').strip(),
                contact_person=request.POST.get('contact_person', ''),
                phone=request.POST.get('phone', ''),
                email=request.POST.get('email') or None,
                address=request.POST.get('address', '')
            )
            messages.success(request, '✓ New supplier added successfully.')
            return redirect('/reports/?tab=suppliers')
        elif action in ('edit_project', 'delete_project'):
            project_obj = get_object_or_404(Project, project_code=request.POST.get('project_code'), is_deleted=False)
            if action == 'delete_project':
                project_obj.is_deleted = True
                project_obj.save(update_fields=['is_deleted'])
                messages.success(request, f"✓ Project [{project_obj.project_code}] removed.")
            else:
                project_obj.project_name = request.POST.get('project_name', project_obj.project_name).strip()
                project_obj.client_name = request.POST.get('client_name', project_obj.client_name).strip()
                project_obj.location = request.POST.get('location', project_obj.location).strip()
                project_obj.status = request.POST.get('status', project_obj.status)
                project_obj.start_date = request.POST.get('start_date') or project_obj.start_date
                project_obj.completion_date = request.POST.get('completion_date') or None
                project_obj.save()
                messages.success(request, f"✓ Project [{project_obj.project_code}] updated.")
            return redirect('/reports/?tab=projects')
        elif action == 'create_project':
            Project.objects.create(
                project_code=request.POST.get('project_code', '').strip(),
                project_name=request.POST.get('project_name', '').strip(),
                client_name=request.POST.get('client_name', '').strip(),
                location=request.POST.get('location', 'Dubai, UAE').strip(),
                status=request.POST.get('status', 'Active'),
                start_date=request.POST.get('start_date') or timezone.now().date(),
                completion_date=request.POST.get('completion_date') or None
            )
            messages.success(request, '✓ New project/site added successfully.')
            return redirect('/reports/?tab=projects')
        elif action == 'create_custom_pr':
            mat_codes = request.POST.getlist('selected_materials')
            pr_code = f"PR-KEN-{timezone.now().strftime('%y%m%d')}-{random.randint(100, 999)}"
            pr_notes = request.POST.get('notes', 'Manual Store Requisition')
            
            items_payload = []
            for code in mat_codes:
                m = Material.objects.filter(material_code=code).first()
                if m:
                    custom_qty = request.POST.get(f'qty_{code}')
                    reorder_val = int(custom_qty) if custom_qty and custom_qty.isdigit() else max(10, m.min_stock_alert * 2)
                    items_payload.append({
                        'code': m.material_code,
                        'desc': m.description,
                        'current_stock': m.current_stock,
                        'unit': m.unit,
                        'reorder_qty': reorder_val
                    })

            PurchaseRequest.objects.create(
                pr_number=pr_code,
                generated_by=f"{request.user.username}, {request.user.get_full_name() or 'Manuja Shehan'}",
                total_items=len(items_payload),
                status='Pending Approval',
                notes=pr_notes,
                items_json=json.dumps(items_payload)
            )
            messages.success(request, f"✓ Created Purchase Request [{pr_code}] for {len(items_payload)} item(s).")
            return redirect('/reports/?tab=purchase')
        elif 'generate_bulk_pr' in request.POST or action == 'auto_bulk_pr':
            low_stock_items = Material.objects.filter(current_stock__lte=F('min_stock_alert'), is_deleted=False)
            count = low_stock_items.count()
            if count == 0:
                messages.error(request, "No items are currently below minimum alert level in Store F20.")
            else:
                pr_code = f"PR-KEN-{timezone.now().strftime('%y%m%d')}-{random.randint(100, 999)}"
                items_payload = []
                for m in low_stock_items:
                    items_payload.append({
                        'code': m.material_code,
                        'desc': m.description,
                        'current_stock': m.current_stock,
                        'unit': m.unit,
                        'reorder_qty': max(50, m.min_stock_alert * 2)
                    })
                PurchaseRequest.objects.create(
                    pr_number=pr_code,
                    generated_by=f"{request.user.username}, {request.user.get_full_name() or 'Manuja Shehan'}",
                    total_items=count,
                    status='Submitted to Procurement',
                    notes=f"Auto-generated Bulk PR for {count} low-stock materials in Store F20.",
                    items_json=json.dumps(items_payload)
                )
                messages.success(request, f"✓ Auto Bulk PR [{pr_code}] generated successfully for {count} item(s)!")
            return redirect('/reports/?tab=purchase')
        elif action == 'create_staff':
            Staff.objects.create(
                staff_no=request.POST.get('staff_no').strip(),
                name=request.POST.get('name').strip(),
                designation=request.POST.get('designation', 'Technician').strip(),
                phone=request.POST.get('phone', ''),
                is_store_custodian=True if request.POST.get('is_store_custodian') == 'on' else False
            )
            messages.success(request, f"✓ Registered Staff [{request.POST.get('staff_no')}].")
            return redirect('/reports/?tab=staff')
        elif action == 'edit_staff':
            st = get_object_or_404(Staff, staff_no=request.POST.get('staff_no'))
            st.name = request.POST.get('name')
            st.designation = request.POST.get('designation')
            st.phone = request.POST.get('phone')
            st.is_store_custodian = True if request.POST.get('is_store_custodian') == 'on' else False
            st.save()
            messages.success(request, f"✓ Updated Staff [{st.staff_no}] details.")
            return redirect('/reports/?tab=staff')
        elif action == 'delete_staff':
            st = get_object_or_404(Staff, staff_no=request.POST.get('staff_no'))
            st.is_deleted = True
            st.save()
            messages.success(request, f"✓ Staff [{st.staff_no}] removed.")
            return redirect('/reports/?tab=staff')

    date_field_tabs = ['toolhistory', 'issues', 'receipts', 'materialreturns', 'purchase', 'deleted', 'gatepassall', 'materialgp', 'toolsgp']
    rows = []
    
    if tab == 'toolhistory':
        qs = ToolTransaction.objects.select_related('tool', 'staff', 'project').order_by('-issue_date')
        if q:
            qs = qs.filter(
                Q(tool__tool_no__icontains=q) |
                Q(tool__tool_name__icontains=q) |
                Q(staff__name__icontains=q) |
                Q(staff__staff_no__icontains=q) |
                Q(project__project_name__icontains=q)
            )
        if state == 'pending':
            qs = qs.filter(status='Issued')
        elif state == 'returned':
            qs = qs.filter(status='Returned')
        if date_from:
            qs = qs.filter(issue_date__date__gte=date_from)
        if date_to:
            qs = qs.filter(issue_date__date__lte=date_to)
        rows = qs

    elif tab == 'toolregister':
        if status == 'Issued':
            pending_qs = ToolTransaction.objects.filter(status='Issued', tool__is_deleted=False).select_related('tool').order_by('-issue_date')
            if q:
                pending_qs = pending_qs.filter(Q(tool__tool_no__icontains=q) | Q(tool__tool_name__icontains=q))
            # Keep this filter consistent with Pending Tools: one row per active issue transaction.
            rows = [tx.tool for tx in pending_qs]
        else:
            qs = Tool.objects.filter(is_deleted=False).order_by('tool_no')
            if q:
                qs = qs.filter(Q(tool_no__icontains=q) | Q(tool_name__icontains=q))
            if status:
                qs = qs.filter(status__iexact=status)
            rows = qs

    elif tab == 'materialstock':
        latest_grn = MaterialTransaction.objects.filter(
            transaction_type='Receive', material_id=OuterRef('pk')
        ).order_by('-date')
        qs = Material.objects.filter(is_deleted=False).annotate(
            low_stock_first=Case(
                When(current_stock__lte=F('min_stock_alert'), then=0),
                default=1,
                output_field=IntegerField(),
            ),
            latest_grn_no=Subquery(latest_grn.values('reference_no')[:1]),
            latest_dn_no=Subquery(latest_grn.values('delivery_note_no')[:1]),
            latest_grn_date=Subquery(latest_grn.values('date')[:1]),
        ).order_by('low_stock_first', 'material_code')
        if q:
            qs = qs.filter(Q(material_code__icontains=q) | Q(description__icontains=q))
        if stock == 'low':
            qs = qs.filter(current_stock__lte=F('min_stock_alert'))
        elif stock == 'ok':
            qs = qs.filter(current_stock__gt=F('min_stock_alert'))
        rows = qs

    elif tab == 'issues':
        qs = MaterialTransaction.objects.filter(transaction_type='Issue').select_related('material', 'project', 'staff').order_by('-date')
        if q:
            qs = qs.filter(Q(material__material_code__icontains=q) | Q(material__description__icontains=q) | Q(staff__name__icontains=q))
        if project_filter:
            qs = qs.filter(project__project_code=project_filter)
        if date_from:
            qs = qs.filter(date__date__gte=date_from)
        if date_to:
            qs = qs.filter(date__date__lte=date_to)
        rows = qs

    elif tab == 'receipts':
        qs = MaterialTransaction.objects.filter(transaction_type='Receive').select_related('material', 'supplier').order_by('-date')
        if q:
            qs = qs.filter(Q(material__material_code__icontains=q) | Q(supplier__company_name__icontains=q) | Q(reference_no__icontains=q) | Q(delivery_note_no__icontains=q) | Q(received_by__icontains=q))
        if date_from:
            qs = qs.filter(date__date__gte=date_from)
        if date_to:
            qs = qs.filter(date__date__lte=date_to)
        rows = qs

    elif tab == 'materialreturns':
        qs = MaterialTransaction.objects.filter(transaction_type='Return').select_related('material', 'project', 'staff').order_by('-date')
        if q:
            qs = qs.filter(Q(material__material_code__icontains=q) | Q(staff__name__icontains=q) | Q(project__project_name__icontains=q))
        if date_from:
            qs = qs.filter(date__date__gte=date_from)
        if date_to:
            qs = qs.filter(date__date__lte=date_to)
        rows = qs

    elif tab == 'purchase':
        qs = PurchaseRequest.objects.all().order_by('-created_date')
        if q:
            qs = qs.filter(Q(pr_number__icontains=q) | Q(generated_by__icontains=q) | Q(notes__icontains=q) | Q(items_json__icontains=q))
        if date_from:
            qs = qs.filter(created_date__date__gte=date_from)
        if date_to:
            qs = qs.filter(created_date__date__lte=date_to)
        rows = _build_pr_item_rows(qs)
        if status:
            status_l = status.lower()
            rows = [r for r in rows if r['item_status'].lower() == status_l]

    elif tab == 'staff':
        qs = Staff.objects.filter(is_deleted=False).order_by('staff_no')
        if q:
            qs = qs.filter(Q(staff_no__icontains=q) | Q(name__icontains=q) | Q(designation__icontains=q))
        rows = qs

    elif tab == 'suppliers':
        qs = Supplier.objects.filter(is_deleted=False).order_by('company_name')
        if q:
            qs = qs.filter(Q(supplier_code__icontains=q) | Q(company_name__icontains=q))
        rows = qs

    elif tab == 'projects':
        qs = Project.objects.filter(is_deleted=False).order_by('-start_date')
        if q:
            qs = qs.filter(Q(project_code__icontains=q) | Q(project_name__icontains=q) | Q(client_name__icontains=q))
        rows = qs

    elif tab in ('gatepassall', 'materialgp', 'toolsgp'):
        qs = GatePass.objects.all().prefetch_related('items').order_by('-created_at')
        if tab == 'materialgp':
            qs = qs.filter(pass_type='Material')
        elif tab == 'toolsgp':
            qs = qs.filter(pass_type='Tool')
        if q:
            qs = qs.filter(
                Q(movement_pass_no__icontains=q) |
                Q(employee_name__icontains=q) |
                Q(employee_id_designation__icontains=q) |
                Q(destination__icontains=q) |
                Q(created_by__icontains=q)
            )
        if date_from:
            qs = qs.filter(created_at__date__gte=date_from)
        if date_to:
            qs = qs.filter(created_at__date__lte=date_to)
        rows = qs

    elif tab == 'deleted':
        deleted_list = []
        for t in Tool.objects.filter(is_deleted=True):
            deleted_list.append({'record_type': 'Tool', 'record_key': t.tool_no, 'label': t.tool_name, 'deleted_at': timezone.now()})
        for m in Material.objects.filter(is_deleted=True):
            deleted_list.append({'record_type': 'Material', 'record_key': m.material_code, 'label': m.description, 'deleted_at': timezone.now()})
        for p in Project.objects.filter(is_deleted=True):
            deleted_list.append({'record_type': 'Project', 'record_key': p.project_code, 'label': p.project_name, 'deleted_at': timezone.now()})
        for s in Supplier.objects.filter(is_deleted=True):
            deleted_list.append({'record_type': 'Supplier', 'record_key': s.supplier_code, 'label': s.company_name, 'deleted_at': timezone.now()})
        for st in Staff.objects.filter(is_deleted=True):
            deleted_list.append({'record_type': 'Staff', 'record_key': st.staff_no, 'label': st.name, 'deleted_at': timezone.now()})
        rows = deleted_list

    filtered_count = len(rows) if isinstance(rows, list) else rows.count()
    paginator = Paginator(rows, 500 if tab in ('issues', 'toolhistory', 'gatepassall', 'materialgp', 'toolsgp') else 25)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    query_params = request.GET.copy()
    query_params.pop('page', None)

    context = {
        'page_title': 'Records & Reports Workspace',
        'page_subtitle': 'Master records, stock control and complete transaction history in one workspace.',
        'page_icon': 'fa-solid fa-file-invoice',
        'page_icon_color': 'from-emerald-600 to-teal-700',
        'tab': tab,
        'q': q,
        'state': state,
        'status': status,
        'stock': stock,
        'project': project_filter,
        'date_from': date_from,
        'date_to': date_to,
        'date_field_tabs': date_field_tabs,
        'projects': Project.objects.filter(is_deleted=False),
        'all_materials': Material.objects.filter(is_deleted=False).order_by('material_code'),
        'rows': page_obj,
        'page_obj': page_obj,
        'filtered_count': filtered_count,
        'filter_active': any([q, state != 'all', status, stock, project_filter, date_from, date_to]),
        'report_querystring': query_params.urlencode(),
        'today': timezone.localdate(),
    }
    return render(request, 'store_app/reports_hub.html', context)

def _next_gate_pass_no():
    year = timezone.localdate().year
    prefix = f'MP/{year}/'
    last = GatePass.objects.filter(movement_pass_no__startswith=prefix).order_by('-id').first()
    if not last:
        return f'{prefix}0001'
    try:
        n = int(last.movement_pass_no.rsplit('/', 1)[1]) + 1
    except (ValueError, IndexError):
        n = GatePass.objects.filter(movement_pass_no__startswith=prefix).count() + 1
    return f'{prefix}{n:04d}'


def _create_gate_pass_from_tool_ids(request, selected_ids):
    if not selected_ids:
        messages.error(request, 'Select at least one issued tool for the Tools Gate Pass.')
        return redirect('tool_pending')
    txs = list(ToolTransaction.objects.filter(id__in=selected_ids, status='Issued').select_related('tool', 'staff', 'project').order_by('id'))
    if len(txs) != len(set(map(int, selected_ids))):
        messages.error(request, 'One or more selected tools are no longer issued and could not be included.')
        return redirect('tool_pending')
    # gate_pass_no is a legacy field populated by the older ERP. Do not use it
    # as proof that a new GatePass exists. A tool transaction is considered
    # already passed only when it is linked to an actual GatePassItem record.
    existing_tool_tx_ids = set(
        GatePassItem.objects.filter(
            tool_transaction_id__in=[tx.id for tx in txs],
            gate_pass__isnull=False,
        ).values_list('tool_transaction_id', flat=True)
    )
    if existing_tool_tx_ids:
        messages.error(request, 'One or more selected tools already have a Gate Pass.')
        return redirect('tool_pending')
    return _create_gate_pass(request, 'Tool', txs=txs)


def _create_gate_pass_from_material_ids(request, selected_ids):
    if not selected_ids:
        messages.error(request, 'Select at least one material issue for the Material Gate Pass.')
        return redirect('/reports/?tab=issues')
    txs = list(MaterialTransaction.objects.filter(id__in=selected_ids, transaction_type='Issue').select_related('material', 'staff', 'project').order_by('id'))
    if len(txs) != len(set(map(int, selected_ids))):
        messages.error(request, 'One or more selected material issues are no longer available.')
        return redirect('/reports/?tab=issues')
    # Check the actual GatePassItem relation, not the legacy gate_pass_no field.
    existing_material_tx_ids = set(
        GatePassItem.objects.filter(material_transaction_id__in=[tx.id for tx in txs])
        .values_list('material_transaction_id', flat=True)
    )
    if existing_material_tx_ids:
        messages.error(request, 'One or more selected material issues already have a Gate Pass.')
        return redirect('/reports/?tab=issues')
    return _create_gate_pass(request, 'Material', material_txs=txs)


@transaction.atomic
def _create_gate_pass(request, pass_type, txs=None, material_txs=None):
    txs = txs or []
    material_txs = material_txs or []
    selected = txs or material_txs
    first = selected[0]
    staff = getattr(first, 'staff', None)
    project = getattr(first, 'project', None)
    staff_key = staff.staff_no if staff else None
    project_key = project.project_code if project else None
    if any((getattr(x, 'staff', None).staff_no if getattr(x, 'staff', None) else None) != staff_key for x in selected):
        messages.error(request, 'Selected records belong to different employees. Create separate Gate Passes for each employee.')
        return redirect('tool_pending' if pass_type == 'Tool' else '/reports/?tab=issues')
    if any((getattr(x, 'project', None).project_code if getattr(x, 'project', None) else None) != project_key for x in selected):
        messages.error(request, 'Selected records belong to different project/site destinations. Create separate Gate Passes for each destination.')
        return redirect('tool_pending' if pass_type == 'Tool' else '/reports/?tab=issues')

    gp = GatePass.objects.create(
        movement_pass_no=_next_gate_pass_no(),
        pass_type=pass_type,
        created_by=f'{request.user.username}, {request.user.get_full_name() or "Manuja Shehan"}',
        employee_name=staff.name if staff else None,
        employee_id_designation=f'{staff.staff_no} / {staff.designation}' if staff else None,
        destination=project.project_name if project else None,
    )
    if pass_type == 'Tool':
        for tx in txs:
            GatePassItem.objects.create(gate_pass=gp, tool_transaction=tx, description=tx.tool.tool_name, quantity='1', serial_asset_no=tx.tool.serial_no or '')
            tx.gate_pass_no = gp.movement_pass_no
            tx.save(update_fields=['gate_pass_no'])
    else:
        for tx in material_txs:
            GatePassItem.objects.create(gate_pass=gp, material_transaction=tx, description=tx.material.description, quantity=f'{tx.quantity} {tx.material.unit.upper()}', serial_asset_no='')
            tx.gate_pass_no = gp.movement_pass_no
            tx.save(update_fields=['gate_pass_no'])
    messages.success(request, f'✓ {gp.get_pass_type_display()} {gp.movement_pass_no} created for {len(selected)} item(s).')
    return redirect('print_gate_pass', gp_id=gp.id)


@login_required
def print_gate_pass(request, gp_id):
    gp = get_object_or_404(GatePass.objects.prefetch_related('items'), id=gp_id)
    pdf = render_gate_pass_pdf(gp)
    response = HttpResponse(pdf.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{gp.movement_pass_no.replace("/", "-")}.pdf"'
    response['X-Frame-Options'] = 'SAMEORIGIN'
    return response


@login_required
def daily_report(request):
    data = collect_daily_data()
    pdf = build_daily_pdf(data)
    response = HttpResponse(pdf.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="Kensova_Daily_Report_{data["day"]:%Y%m%d}.pdf"'
    return response


@login_required
def send_daily_report_now(request):
    if request.method != 'POST':
        return redirect('reports_hub')
    data = collect_daily_data()
    pdf = build_daily_pdf(data)
    xlsx = build_daily_excel(data)
    try:
        send_via_outlook(pdf, xlsx, data['day'])
        messages.success(request, f'✓ Daily management report sent from Outlook to {len(TO_EMAILS)} To recipient(s) with CC recipients.')
    except Exception as exc:
        messages.error(request, f'Could not send daily report from Outlook: {exc}')
    return redirect('dashboard')



def _ai_history_db():
    """Use a small independent SQLite file so Ask AI never alters or blocks the ERP database."""
    from django.conf import settings
    path = os.path.join(str(settings.BASE_DIR), 'ai_chat_history.sqlite3')
    conn = sqlite3.connect(path, timeout=5)
    conn.execute('PRAGMA journal_mode=WAL')
    conn.execute("""CREATE TABLE IF NOT EXISTS chat_messages (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER NOT NULL,
        username TEXT NOT NULL,
        role TEXT NOT NULL,
        content TEXT NOT NULL,
        action TEXT NOT NULL DEFAULT 'chat',
        created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    )""")
    conn.commit()
    return conn

@login_required
def ai_history(request):
    try:
        conn = _ai_history_db()
        rows = conn.execute(
            'SELECT role, content, action, created_at FROM chat_messages WHERE user_id=? ORDER BY id ASC',
            (request.user.pk,)
        ).fetchall()
        conn.close()
        return JsonResponse({'messages':[{'role':r[0], 'content':r[1], 'action':r[2], 'created_at':r[3]} for r in rows]})
    except Exception:
        return JsonResponse({'messages':[]})

def _gemini_error_message(exc):
    try:
        raw = exc.read().decode('utf-8', errors='replace')
        payload = json.loads(raw)
        api_error = payload.get('error') or {}
        return str(api_error.get('message') or '').strip()
    except Exception:
        return ''


def _gemini_generate(api_key, model_id, system, recent, message):
    """Send one text/chat request using Google's Gemini Interactions API."""
    transcript = []
    for role, content in recent:
        label = 'Assistant' if role == 'assistant' else 'User'
        transcript.append(f'{label}: {content}')
    transcript.append(f'User: {message}')

    data = json.dumps({
        'model': model_id,
        'input': '\n\n'.join(transcript),
        'system_instruction': system,
        'store': False,
        'generation_config': {
            'max_output_tokens': 1800,
            'thinking_level': 'minimal',
        },
    }).encode('utf-8')

    req = urllib.request.Request(
        'https://generativelanguage.googleapis.com/v1beta/interactions',
        data=data,
        headers={
            'x-goog-api-key': api_key,
            'Content-Type': 'application/json',
        },
        method='POST',
    )
    with urllib.request.urlopen(req, timeout=35) as resp:
        result = json.loads(resp.read().decode('utf-8'))

    if result.get('status') == 'failed':
        err = result.get('error') or {}
        raise RuntimeError(str(err.get('message') or 'The AI interaction failed.'))

    texts = []
    for step in result.get('steps') or []:
        if step.get('type') != 'model_output':
            continue
        for item in step.get('content') or []:
            if item.get('type') == 'text' and item.get('text'):
                texts.append(str(item['text']))
    return '\n'.join(texts).strip()


@login_required
@require_POST
def ai_chat(request):
    try:
        payload=json.loads(request.body.decode('utf-8') or '{}')
    except Exception:
        return JsonResponse({'error':'Invalid request body.'},status=400)
    message=str(payload.get('message') or '').strip()
    action=str(payload.get('action') or 'chat').strip().lower()
    if not message:
        return JsonResponse({'error':'Please enter some text.'},status=400)
    if len(message)>12000:
        return JsonResponse({'error':'Text is too long. Please keep it under 12,000 characters.'},status=400)

    api_key=os.environ.get('GEMINI_API_KEY','').strip() or os.environ.get('GOOGLE_API_KEY','').strip()
    if not api_key:
        return JsonResponse({'error':'AI service key is not configured on this computer.'},status=503)

    system={'correct':'Correct the user text into clear professional English. Preserve meaning, names, IDs, numbers, and formatting. Return only the corrected text.',
            'sinhala':'Translate the user text accurately into Sinhala (Sri Lankan language). Preserve names, IDs, numbers, and formatting. Return only the Sinhala translation.',
            'chat':'You are the Kensova ERP assistant. Be concise, practical, and professional. Help with writing, translation, and operations questions.'}.get(action,'You are the Kensova ERP assistant. Be concise and professional.')

    conn=None
    try:
        conn=_ai_history_db()
        recent=conn.execute(
            'SELECT role, content FROM chat_messages WHERE user_id=? ORDER BY id DESC LIMIT 20',
            (request.user.pk,)
        ).fetchall()
        recent=list(reversed(recent))
    except Exception:
        recent=[]

    try:
        if conn is None: conn=_ai_history_db()
        conn.execute('INSERT INTO chat_messages(user_id,username,role,content,action) VALUES(?,?,?,?,?)',(request.user.pk,request.user.username,'user',message,action))
        conn.commit()
    except Exception:
        pass

    # Use current Gemini Interactions API models. Flash-Lite is preferred for low-latency office tasks;
    # Gemini 3.6 Flash is the fallback if the first model is unavailable for the project/key.
    models = ('gemini-3.5-flash-lite', 'gemini-3.6-flash')
    reply = ''
    selected_model = ''
    last_code = None
    last_api_message = ''

    for model_id in models:
        try:
            reply = _gemini_generate(api_key, model_id, system, recent, message)
            if reply:
                selected_model = model_id
                break
            last_api_message = 'The AI service returned an empty response.'
        except urllib.error.HTTPError as exc:
            last_code = exc.code
            last_api_message = _gemini_error_message(exc)
            # Invalid/revoked authentication should not be retried with another model.
            if exc.code in (401,):
                break
            # Model/quota/project-specific errors may succeed with the fallback model.
            if exc.code in (400, 403, 404, 429):
                continue
            break
        except Exception as exc:
            last_api_message = str(exc)
            break

    if not reply:
        if conn:
            try: conn.close()
            except Exception: pass
        if last_code == 401:
            msg = 'AI API key was rejected. The configured key is invalid or revoked.'
        elif last_code == 429:
            msg = 'AI request limit reached. Please wait a short time and try again.'
        elif last_api_message:
            msg = 'AI service error: ' + last_api_message[:350]
        else:
            msg = 'AI service returned an error. Please try again.'
        return JsonResponse({'error':msg},status=502)

    try:
        if conn is None: conn=_ai_history_db()
        conn.execute('INSERT INTO chat_messages(user_id,username,role,content,action) VALUES(?,?,?,?,?)',(request.user.pk,request.user.username,'assistant',reply,action))
        conn.commit()
        conn.close()
    except Exception:
        if conn:
            try: conn.close()
            except Exception: pass
    return JsonResponse({'reply':reply,'model':selected_model})
